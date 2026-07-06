#!/usr/bin/env python3
"""Parse lab documents into Markdown and optionally upload them to the RAG API."""

from __future__ import annotations

import argparse
import json
import mimetypes
import re
import shutil
import subprocess
import sys
import tempfile
import traceback
import urllib.error
import urllib.request
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Callable


SUPPORTED_EXTENSIONS = {".pdf", ".pptx", ".docx", ".md", ".markdown", ".txt"}


@dataclass
class DocumentMetadata:
    title: str
    doc_type: str
    topic: str
    owner: str
    owner_status: str
    visibility: str
    source_path: str
    parser: str
    parse_status: str
    parse_error: str = ""
    mime_type: str = ""
    quality_status: str = "unchecked"
    quality_score: int = 0
    quality_notes: str = ""


def read_text(path: Path) -> str:
    for encoding in ("utf-8", "utf-8-sig", "gb18030"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return path.read_text(errors="replace")


def parse_with_docling(path: Path) -> str:
    try:
        from docling.document_converter import DocumentConverter
    except ImportError as exc:
        raise RuntimeError("docling is not installed; install with `pip install docling`") from exc

    converter = DocumentConverter()
    result = converter.convert(str(path))
    return result.document.export_to_markdown()


def parse_with_marker(path: Path, timeout_seconds: int) -> str:
    marker_single = shutil.which("marker_single")
    if not marker_single:
        candidate = Path(sys.executable).parent / "marker_single"
        if candidate.exists():
            marker_single = str(candidate)
    if not marker_single:
        raise RuntimeError("marker_single is not installed; install with `pip install marker-pdf`")

    with tempfile.TemporaryDirectory(prefix="marker-out-") as tmp:
        out_dir = Path(tmp)
        cmd = [
            marker_single,
            str(path),
            "--output_format",
            "markdown",
            "--output_dir",
            str(out_dir),
            "--disable_image_extraction",
        ]
        try:
            result = subprocess.run(
                cmd,
                text=True,
                capture_output=True,
                check=False,
                timeout=timeout_seconds,
            )
        except subprocess.TimeoutExpired as exc:
            raise RuntimeError(f"marker_single timed out after {timeout_seconds}s") from exc
        if result.returncode != 0:
            stderr = result.stderr.strip() or result.stdout.strip()
            raise RuntimeError(f"marker_single failed: {stderr}")

        markdown_files = sorted(out_dir.rglob("*.md"))
        if not markdown_files:
            raise RuntimeError("marker_single did not produce a Markdown file")
        return read_text(markdown_files[0])


def parse_docx_simple(path: Path) -> str:
    try:
        import docx
    except ImportError as exc:
        raise RuntimeError("python-docx is not installed; install with `pip install python-docx`") from exc

    document = docx.Document(str(path))
    lines: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = paragraph.style.name.lower() if paragraph.style else ""
        if "heading 1" in style_name or style_name == "title":
            lines.append(f"# {text}")
        elif "heading 2" in style_name:
            lines.append(f"## {text}")
        elif "heading 3" in style_name:
            lines.append(f"### {text}")
        else:
            lines.append(text)
    return "\n\n".join(lines)


def parse_pptx_simple(path: Path) -> str:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise RuntimeError("python-pptx is not installed; install with `pip install python-pptx`") from exc

    presentation = Presentation(str(path))
    sections: list[str] = []
    for idx, slide in enumerate(presentation.slides, start=1):
        texts: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text = shape.text.strip()
                if text:
                    texts.append(text)
        if texts:
            sections.append(f"## Slide {idx}\n\n" + "\n\n".join(texts))
    return "\n\n".join(sections)


def parse_pdf_simple(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("pypdf is not installed; install with `pip install pypdf`") from exc

    reader = PdfReader(str(path))
    pages: list[str] = []
    for idx, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(f"## Page {idx}\n\n{text}")
    return "\n\n".join(pages)


def parse_simple(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".md", ".markdown", ".txt"}:
        return read_text(path)
    if suffix == ".docx":
        return parse_docx_simple(path)
    if suffix == ".pptx":
        return parse_pptx_simple(path)
    if suffix == ".pdf":
        return parse_pdf_simple(path)
    raise RuntimeError(f"unsupported extension: {suffix}")


def parse_document(path: Path, parser: str, parser_timeout: int) -> tuple[str, str]:
    if path.suffix.lower() in {".md", ".markdown", ".txt"}:
        return read_text(path), "text"
    if parser == "marker":
        return parse_with_marker(path, parser_timeout), "marker"
    if parser == "docling":
        return parse_with_docling(path), "docling"
    if parser == "simple":
        return parse_simple(path), "simple"

    if path.suffix.lower() == ".pdf" and infer_doc_type(path) == "paper":
        try:
            return parse_with_marker(path, parser_timeout), "marker"
        except Exception:
            try:
                return parse_with_docling(path), "docling"
            except Exception:
                return parse_simple(path), "simple"

    try:
        return parse_with_docling(path), "docling"
    except Exception:
        return parse_simple(path), "simple"


def clean_paper_markdown(markdown: str) -> str:
    replacements = {
        "ﬁ": "fi",
        "ﬂ": "fl",
        "ﬀ": "ff",
        "ﬃ": "ffi",
        "ﬄ": "ffl",
        "\u00ad": "",
    }
    for src, dst in replacements.items():
        markdown = markdown.replace(src, dst)

    markdown = markdown.replace("\r\n", "\n").replace("\r", "\n")
    markdown = re.sub(r"([A-Za-z])-\n([a-z])", r"\1\2", markdown)

    section_names = {
        "ABSTRACT",
        "INTRODUCTION",
        "BACKGROUND",
        "RELATED WORK",
        "DESIGN",
        "IMPLEMENTATION",
        "EVALUATION",
        "DISCUSSION",
        "CONCLUSION",
        "CONCLUSIONS",
        "REFERENCES",
        "ACKNOWLEDGMENTS",
        "ACKNOWLEDGEMENTS",
    }

    cleaned: list[str] = []
    skip_permission_block = False
    for raw_line in markdown.splitlines():
        line = raw_line.strip()
        if re.search(r"Permission to make digital or hard copies", line, flags=re.IGNORECASE):
            skip_permission_block = True
            continue
        if skip_permission_block:
            if re.search(r"DOI:|ISBN|ACM", line, flags=re.IGNORECASE):
                continue
            if re.match(r"^[A-Z][a-z].{20,}$", line):
                skip_permission_block = False
            else:
                continue
        if not line:
            if cleaned and cleaned[-1] != "":
                cleaned.append("")
            continue
        if re.fullmatch(r"\d{1,4}", line):
            continue
        if re.fullmatch(r"#+\s*Page\s+\d+", line, flags=re.IGNORECASE):
            continue
        if re.search(r"Copyrights? for components", line, flags=re.IGNORECASE):
            continue
        if re.search(r"Request permissions from", line, flags=re.IGNORECASE):
            continue

        heading_candidate = re.sub(r"^#+\s*", "", line).strip()
        heading_key = re.sub(r"^[IVXLC]+\.?\s+", "", heading_candidate, flags=re.IGNORECASE)
        heading_key = re.sub(r"^\d+(?:\.\d+)*\.?\s+", "", heading_key).strip().upper()
        if heading_key in section_names:
            cleaned.append("")
            cleaned.append(f"## {heading_key.title()}")
            cleaned.append("")
            continue
        if re.fullmatch(r"(?:[IVXLC]+|\d+(?:\.\d+)*)\.?\s+[A-Z][A-Z0-9 ,:()/\\-]{3,}", heading_candidate):
            title = re.sub(r"^(?:[IVXLC]+|\d+(?:\.\d+)*)\.?\s+", "", heading_candidate).title()
            cleaned.append("")
            cleaned.append(f"## {title}")
            cleaned.append("")
            continue

        cleaned.append(line)

    text = "\n".join(cleaned)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def quality_report(markdown: str, doc_type: str, parser_used: str) -> tuple[str, int, str]:
    text = markdown.strip()
    lines = text.splitlines()
    bad_chars = sum(text.count(ch) for ch in ["�", "□", "\ufffd"])
    long_lines = sum(1 for line in lines if len(line) > 240)
    headings = [
        line for line in lines
        if line.startswith("#") and not re.match(r"#+\s*Page\s+\d+", line, flags=re.IGNORECASE)
    ]
    real_sections = [
        line for line in headings
        if not re.match(r"#\s+[^#]+$", line) and "<!--" not in line
    ]

    score = 100
    notes: list[str] = []
    if len(text) < 3000:
        score -= 30
        notes.append("too_short")
    if doc_type == "paper":
        if len(text) < 12000:
            score -= 25
            notes.append("paper_too_short")
        if len(real_sections) < 3:
            score -= 25
            notes.append("few_sections")
        if parser_used == "simple":
            score -= 20
            notes.append("fallback_parser")
            score = min(score, 60)
    if bad_chars > 20:
        score -= 20
        notes.append("bad_chars")
    if long_lines > 10:
        score -= 10
        notes.append("many_long_lines")

    score = max(score, 0)
    if doc_type == "paper" and score < 70:
        return "low_quality", score, ",".join(notes)
    if score < 50:
        return "low_quality", score, ",".join(notes)
    return "ok", score, ",".join(notes) or "ok"


def infer_doc_type(path: Path) -> str:
    lower = str(path).lower()
    name = path.name.lower()
    if "paper" in lower or path.parent.name.lower() in {"papers", "paper_pdf"}:
        return "paper"
    if "开题" in name or "proposal" in lower:
        return "proposal"
    if "综述" in name or "survey" in lower:
        return "survey"
    if path.suffix.lower() == ".pptx" or "专题" in name or "seminar" in lower:
        return "seminar"
    if "thesis" in lower or "毕业" in name:
        return "thesis"
    return "document"


def infer_owner(path: Path, root: Path, default_owner: str) -> str:
    rel = path.relative_to(root)
    if len(rel.parts) >= 2 and rel.parts[0] not in {"paper_pdf", "papers", "projects"}:
        return rel.parts[0]
    return default_owner


def infer_topic(path: Path, root: Path, default_topic: str) -> str:
    rel = path.relative_to(root)
    if len(rel.parts) >= 2 and rel.parts[0] in {"paper_pdf", "papers"}:
        return rel.parts[1] if rel.parts[0] == "papers" and len(rel.parts) >= 3 else default_topic
    return default_topic


def markdown_with_header(markdown: str, metadata: DocumentMetadata) -> str:
    body = markdown.strip()
    header = [
        f"# {metadata.title}",
        "",
        "<!--",
        f"doc_type: {metadata.doc_type}",
        f"topic: {metadata.topic}",
        f"owner: {metadata.owner}",
        f"owner_status: {metadata.owner_status}",
        f"source_path: {metadata.source_path}",
        f"parser: {metadata.parser}",
        f"quality_status: {metadata.quality_status}",
        f"quality_score: {metadata.quality_score}",
        f"quality_notes: {metadata.quality_notes}",
        "-->",
        "",
    ]
    if body.startswith("# "):
        return "\n".join(header[2:]) + body + "\n"
    return "\n".join(header) + body + "\n"


def login(api_base: str, username: str, password: str) -> str:
    payload = json.dumps({"username": username, "password": password}).encode("utf-8")
    req = urllib.request.Request(
        f"{api_base.rstrip('/')}/api/auth/login",
        data=payload,
        headers={"Content-Type": "application/json"},
        method="POST",
    )
    with urllib.request.urlopen(req, timeout=30) as resp:
        data = json.loads(resp.read().decode("utf-8"))
    return data["token"]


def upload_document(api_base: str, token: str, filename: str, content: str, visibility: str) -> dict:
    payload = json.dumps(
        {"filename": filename, "content": content, "visibility": visibility},
        ensure_ascii=False,
    ).encode("utf-8")
    req = urllib.request.Request(
        f"{api_base.rstrip('/')}/api/documents",
        data=payload,
        headers={"Content-Type": "application/json", "Authorization": f"Bearer {token}"},
        method="POST",
    )
    with urllib.request.urlopen(req, timeout=60) as resp:
        return json.loads(resp.read().decode("utf-8"))


def iter_documents(root: Path, only: str = "") -> list[Path]:
    files = sorted(
        path
        for path in root.rglob("*")
        if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS
    )
    if not only:
        return files
    return [path for path in files if only in path.relative_to(root).as_posix()]


def write_jsonl(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as fh:
        for row in rows:
            fh.write(json.dumps(row, ensure_ascii=False) + "\n")


def process_file(
    path: Path,
    root: Path,
    output_root: Path,
    args: argparse.Namespace,
    upload_fn: Callable[[str, str, str], dict] | None,
) -> dict:
    rel = path.relative_to(root)
    output_md = output_root / rel.with_suffix(".md")
    output_meta = output_md.with_suffix(".metadata.json")
    metadata = DocumentMetadata(
        title=path.stem,
        doc_type=args.doc_type or infer_doc_type(path),
        topic=args.topic or infer_topic(path, root, root.name),
        owner=args.owner or infer_owner(path, root, "lab"),
        owner_status=args.owner_status,
        visibility=args.visibility,
        source_path=str(path),
        parser=args.parser,
        parse_status="processing",
        mime_type=mimetypes.guess_type(path.name)[0] or "",
    )

    row = {"source_path": str(path), "output_path": str(output_md), "uploaded": False}
    try:
        if output_md.exists() and not args.force:
            parser_used = "cached"
            if output_meta.exists():
                try:
                    cached_metadata = json.loads(read_text(output_meta))
                    parser_used = cached_metadata.get("parser", parser_used)
                except (json.JSONDecodeError, OSError):
                    pass
            metadata.parser = parser_used
        else:
            markdown, parser_used = parse_document(path, args.parser, args.parser_timeout)
            if not markdown.strip():
                raise RuntimeError("parser returned empty content")
            if metadata.doc_type == "paper":
                markdown = clean_paper_markdown(markdown)
            metadata.parser = parser_used
            quality_status, quality_score, quality_notes = quality_report(
                markdown, metadata.doc_type, parser_used)
            metadata.quality_status = quality_status
            metadata.quality_score = quality_score
            metadata.quality_notes = quality_notes
            metadata.parse_status = "ready" if quality_status == "ok" else "low_quality"
            output_md.parent.mkdir(parents=True, exist_ok=True)
            output_md.write_text(markdown_with_header(markdown, metadata), encoding="utf-8")

        parsed_content = read_text(output_md)
        if output_meta.exists() and output_md.exists() and not args.force:
            try:
                cached_metadata = json.loads(read_text(output_meta))
                metadata.parse_status = cached_metadata.get("parse_status", "ready")
                metadata.quality_status = cached_metadata.get("quality_status", "unchecked")
                metadata.quality_score = int(cached_metadata.get("quality_score", 0))
                metadata.quality_notes = cached_metadata.get("quality_notes", "")
            except (json.JSONDecodeError, OSError, ValueError):
                pass
        if metadata.quality_status == "unchecked":
            quality_status, quality_score, quality_notes = quality_report(
                parsed_content, metadata.doc_type, metadata.parser)
            metadata.quality_status = quality_status
            metadata.quality_score = quality_score
            metadata.quality_notes = quality_notes
            metadata.parse_status = "ready" if quality_status == "ok" else "low_quality"
        output_meta.write_text(
            json.dumps(asdict(metadata), ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )

        row.update(asdict(metadata))
        if upload_fn:
            if metadata.parse_status != "ready":
                row["upload_skipped"] = metadata.parse_status
            else:
                upload_name = rel.with_suffix(".md").as_posix()
                upload_result = upload_fn(upload_name, parsed_content, metadata.visibility)
                row["uploaded"] = True
                row["upload_result"] = upload_result
    except Exception as exc:
        metadata.parse_status = "failed"
        metadata.parse_error = str(exc)
        output_meta.parent.mkdir(parents=True, exist_ok=True)
        output_meta.write_text(
            json.dumps(asdict(metadata), ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        row.update(asdict(metadata))
        if args.verbose:
            traceback.print_exc()
    return row


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Parse PDF/PPTX/DOCX lab documents into Markdown for the RAG system."
    )
    parser.add_argument("--input", default="data/tenants/RDMA", help="Source document directory")
    parser.add_argument("--output", default="data/parsed/RDMA", help="Parsed output directory")
    parser.add_argument("--parser", choices=["auto", "marker", "docling", "simple"], default="auto")
    parser.add_argument("--visibility", choices=["group", "public"], default="group")
    parser.add_argument("--topic", default="", help="Override topic metadata for every file")
    parser.add_argument("--owner", default="", help="Override owner metadata for every file")
    parser.add_argument("--owner-status", default="active", choices=["active", "alumni"])
    parser.add_argument("--doc-type", default="", help="Override doc_type metadata for every file")
    parser.add_argument("--only", default="", help="Only process files whose relative path contains this text")
    parser.add_argument("--force", action="store_true", help="Re-parse files with existing Markdown output")
    parser.add_argument("--parser-timeout", type=int, default=180,
                        help="Timeout in seconds for external parsers such as marker")
    parser.add_argument("--upload", action="store_true", help="Upload parsed Markdown to the backend")
    parser.add_argument("--api-base", default="http://localhost:9006")
    parser.add_argument("--token", default="", help="JWT token for --upload")
    parser.add_argument("--username", default="", help="Login username for --upload when --token is omitted")
    parser.add_argument("--password", default="", help="Login password for --upload when --token is omitted")
    parser.add_argument("--verbose", action="store_true")
    return parser


def main() -> int:
    args = build_parser().parse_args()
    input_root = Path(args.input).resolve()
    output_root = Path(args.output).resolve()

    if not input_root.exists():
        print(f"Input directory does not exist: {input_root}", file=sys.stderr)
        return 2

    token = args.token
    if args.upload and not token:
        if not args.username or not args.password:
            print("--upload requires --token or --username/--password", file=sys.stderr)
            return 2
        try:
            token = login(args.api_base, args.username, args.password)
        except (urllib.error.URLError, KeyError, json.JSONDecodeError) as exc:
            print(f"Login failed: {exc}", file=sys.stderr)
            return 1

    upload_fn = None
    if args.upload:
        upload_fn = lambda filename, content, visibility: upload_document(
            args.api_base, token, filename, content, visibility
        )

    files = iter_documents(input_root, args.only)
    if not files:
        print("No supported documents matched the input filters", file=sys.stderr)
        return 2

    rows: list[dict] = []
    for path in files:
        row = process_file(path, input_root, output_root, args, upload_fn)
        rows.append(row)
        status = row.get("parse_status", "unknown")
        print(f"{status:7s} {path.relative_to(input_root)}")

    write_jsonl(output_root / "metadata.jsonl", rows)
    failures = [row for row in rows if row.get("parse_status") not in {"ready", "low_quality"}]
    low_quality = [row for row in rows if row.get("parse_status") == "low_quality"]
    error_path = output_root / "parse_errors.jsonl"
    quality_path = output_root / "quality_report.jsonl"
    if failures:
        write_jsonl(error_path, failures)
    elif error_path.exists():
        error_path.unlink()
    write_jsonl(quality_path, rows)

    ready_count = sum(1 for row in rows if row.get("parse_status") == "ready")
    print(f"\nProcessed {len(rows)} files: {ready_count} ready, {len(low_quality)} low_quality, {len(failures)} failed")
    print(f"Metadata: {output_root / 'metadata.jsonl'}")
    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main())
