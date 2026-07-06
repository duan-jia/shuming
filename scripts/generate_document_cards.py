#!/usr/bin/env python3
"""Generate lightweight document cards for lab RAG discovery."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import sys
import traceback
import urllib.error
import urllib.request
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any


DEFAULT_ENDPOINT = "https://api.deepseek.com/v1/chat/completions"
SUPPORTED_EXTENSIONS = {".pdf", ".pptx", ".docx", ".md", ".markdown", ".txt"}


@dataclass
class DocumentCard:
    schema_version: str
    title: str
    aliases: list[str]
    doc_type: str
    owner: str
    topic: str
    year: int | None
    keywords: list[str]
    summary: str
    main_contribution: str
    useful_for: list[str]
    questions_answered: list[str]
    source_path: str


def clean_text(text: str) -> str:
    replacements = {
        "ﬁ": "fi",
        "ﬂ": "fl",
        "ﬀ": "ff",
        "ﬃ": "ffi",
        "ﬄ": "ffl",
        "\u00ad": "",
    }
    for src, dst in replacements.items():
        text = text.replace(src, dst)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"([A-Za-z])-\n([a-z])", r"\1\2", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def read_text(path: Path) -> str:
    for encoding in ("utf-8", "utf-8-sig", "gb18030"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return path.read_text(errors="replace")


def read_pdf_pages(path: Path, first_pages: int, last_pages: int) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("pypdf is not installed; install with `pip install pypdf`") from exc

    reader = PdfReader(str(path))
    page_count = len(reader.pages)
    page_indexes = set(range(min(first_pages, page_count)))
    if last_pages > 0:
        page_indexes.update(range(max(0, page_count - last_pages), page_count))

    pages: list[str] = []
    for idx in sorted(page_indexes):
        text = reader.pages[idx].extract_text() or ""
        if text.strip():
            pages.append(f"--- Page {idx + 1} ---\n{text.strip()}")
    return clean_text("\n\n".join(pages))


def read_pptx_text(path: Path, max_slides: int) -> str:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise RuntimeError("python-pptx is not installed; install with `pip install python-pptx`") from exc

    presentation = Presentation(str(path))
    sections: list[str] = []
    for idx, slide in enumerate(presentation.slides, start=1):
        if idx > max_slides:
            break
        texts: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text = shape.text.strip()
                if text:
                    texts.append(text)
        if texts:
            sections.append(f"--- Slide {idx} ---\n" + "\n".join(texts))
    return clean_text("\n\n".join(sections))


def read_docx_text(path: Path, max_paragraphs: int) -> str:
    try:
        import docx
    except ImportError as exc:
        raise RuntimeError("python-docx is not installed; install with `pip install python-docx`") from exc

    document = docx.Document(str(path))
    paragraphs: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            paragraphs.append(text)
        if len(paragraphs) >= max_paragraphs:
            break
    return clean_text("\n\n".join(paragraphs))


def infer_doc_type(path: Path) -> str:
    lower = path.as_posix().lower()
    name = path.name.lower()
    if "paper" in lower or path.parent.name.lower() in {"papers", "paper_pdf"}:
        return "paper"
    if "开题" in name or "proposal" in lower:
        return "proposal"
    if "综述" in name or "survey" in lower:
        return "survey"
    if path.suffix.lower() == ".pptx" or "专题" in name or "seminar" in lower:
        return "slides"
    if "thesis" in lower or "毕业" in name:
        return "thesis"
    return "document"


def safe_slug(value: str) -> str:
    slug = re.sub(r"[^A-Za-z0-9._-]+", "_", value).strip("._-").lower()
    return slug or "document"


def safe_output_slug(rel: Path) -> str:
    rel_text = rel.with_suffix("").as_posix()
    digest = hashlib.sha1(rel_text.encode("utf-8")).hexdigest()[:10]
    return f"{safe_slug(rel_text)}_{digest}"


def infer_owner(path: Path, root: Path, default_owner: str) -> str:
    rel = path.relative_to(root)
    if len(rel.parts) >= 2 and rel.parts[0] not in {"paper_pdf", "papers"}:
        return safe_slug(rel.parts[0])
    return default_owner


def build_model_context(path: Path, args: argparse.Namespace) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        text = read_pdf_pages(path, args.first_pages, args.last_pages)
    elif suffix == ".pptx":
        text = read_pptx_text(path, args.max_slides)
    elif suffix == ".docx":
        text = read_docx_text(path, args.max_paragraphs)
    elif suffix in {".md", ".markdown", ".txt"}:
        text = clean_text(read_text(path)[: args.max_chars])
    else:
        raise RuntimeError(f"unsupported extension: {suffix}")

    if not text:
        raise RuntimeError("text extraction returned empty content")
    return (
        f"Filename: {path.name}\n"
        f"Inferred doc_type: {infer_doc_type(path)}\n\n"
        f"Extracted text sample:\n{text[: args.max_chars]}"
    )


def prompt_for_card(context: str, source_path: str, doc_type: str, owner: str, topic: str) -> list[dict[str, str]]:
    schema = {
        "schema_version": "document_card.v2",
        "title": "",
        "aliases": [],
        "doc_type": doc_type,
        "owner": owner,
        "topic": topic,
        "year": None,
        "keywords": [],
        "summary": "",
        "main_contribution": "",
        "useful_for": [],
        "questions_answered": [],
        "source_path": source_path,
    }
    return [
        {
            "role": "system",
            "content": (
                "You extract concise metadata for a lab RAG discovery system. "
                "Return strict JSON only. Do not include markdown fences or commentary. "
                "Write title, keywords, summary, main_contribution, useful_for, and questions_answered in English. "
                "Preserve Chinese names, original Chinese titles, romanized names, acronyms, and source filename terms in aliases."
            ),
        },
        {
            "role": "user",
            "content": (
                "Create one document card from the extracted text. The card is used for retrieval, "
                "so prefer factual keywords and question-like phrases users may ask. "
                "Keep summary within 100 English words. "
                "Use 3-8 keywords, 2-5 useful_for items, and 3-6 questions_answered items. "
                "All semantic fields must be English except aliases. "
                "aliases must include useful Chinese names/titles and romanized variants when present, "
                "for example 段佳兵, Duan Jiabing, duanjiabing. "
                "If year is unknown, use null. Preserve the provided doc_type, owner, topic, and source_path unless clearly wrong.\n\n"
                f"Required JSON schema:\n{json.dumps(schema, ensure_ascii=False)}\n\n"
                f"Extracted text:\n{context}"
            ),
        },
    ]


def call_chat_completion(endpoint: str, api_key: str, model: str, messages: list[dict[str, str]],
                         temperature: float, timeout: int) -> str:
    payload = json.dumps(
        {
            "model": model,
            "messages": messages,
            "temperature": temperature,
            "response_format": {"type": "json_object"},
        },
        ensure_ascii=False,
    ).encode("utf-8")
    req = urllib.request.Request(
        endpoint,
        data=payload,
        headers={
            "Content-Type": "application/json",
            "Authorization": f"Bearer {api_key}",
        },
        method="POST",
    )
    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            data = json.loads(resp.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        body = exc.read().decode("utf-8", errors="replace").strip()
        detail = body[:1000] if body else exc.reason
        raise RuntimeError(f"HTTP {exc.code}: {detail}") from exc
    return data["choices"][0]["message"]["content"]


def parse_json_object(text: str) -> dict[str, Any]:
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", text, flags=re.DOTALL)
        if not match:
            raise
        return json.loads(match.group(0))


def string_list(value: Any) -> list[str]:
    if isinstance(value, list):
        return [str(item).strip() for item in value if str(item).strip()]
    if isinstance(value, str) and value.strip():
        return [value.strip()]
    return []


def coerce_card(raw: dict[str, Any], source_path: Path, doc_type: str, owner: str, topic: str) -> DocumentCard:
    year_value = raw.get("year")
    year: int | None = None
    if isinstance(year_value, int):
        year = year_value
    elif isinstance(year_value, str) and re.fullmatch(r"\d{4}", year_value.strip()):
        year = int(year_value.strip())

    return DocumentCard(
        schema_version="document_card.v2",
        title=str(raw.get("title") or source_path.stem).strip(),
        aliases=merge_aliases(string_list(raw.get("aliases")), source_path, owner),
        doc_type=str(raw.get("doc_type") or doc_type).strip() or doc_type,
        owner=str(raw.get("owner") or owner).strip() or owner,
        topic=str(raw.get("topic") or topic).strip() or topic,
        year=year,
        keywords=string_list(raw.get("keywords")),
        summary=str(raw.get("summary") or "").strip(),
        main_contribution=str(raw.get("main_contribution") or "").strip(),
        useful_for=string_list(raw.get("useful_for")),
        questions_answered=string_list(raw.get("questions_answered")),
        source_path=str(source_path),
    )


def first_content_lines(context: str, max_chars: int) -> str:
    marker = "Extracted text sample:\n"
    text = context.split(marker, 1)[1] if marker in context else context
    lines = [line.strip() for line in text.splitlines() if line.strip() and not line.startswith("--- ")]
    return clean_text(" ".join(lines))[:max_chars]


def contains_cjk(value: str) -> bool:
    return any("\u4e00" <= ch <= "\u9fff" for ch in value)


def card_needs_english_normalization(card: DocumentCard) -> bool:
    values: list[str] = [
        card.title,
        card.summary,
        card.main_contribution,
        *card.keywords,
        *card.useful_for,
        *card.questions_answered,
    ]
    return any(contains_cjk(value) for value in values)


def normalize_card_to_english(card: DocumentCard, endpoint: str, api_key: str, model: str,
                              temperature: float, timeout: int) -> DocumentCard:
    messages = [
        {
            "role": "system",
            "content": (
                "Normalize a document card for an English RAG index. Return strict JSON only. "
                "Translate title, keywords, summary, main_contribution, useful_for, and questions_answered into English. "
                "Do not translate aliases, owner, topic, source_path, doc_type, year, or schema_version. "
                "Preserve all aliases exactly, especially Chinese names and Chinese titles."
            ),
        },
        {
            "role": "user",
            "content": json.dumps(asdict(card), ensure_ascii=False),
        },
    ]
    raw = parse_json_object(call_chat_completion(endpoint, api_key, model, messages, temperature, timeout))
    normalized = coerce_card(raw, Path(card.source_path), card.doc_type, card.owner, card.topic)
    if not normalized.aliases:
        normalized.aliases = card.aliases
    return normalized


def merge_aliases(raw_aliases: list[str], source_path: Path, owner: str) -> list[str]:
    aliases: list[str] = []

    def add(value: str) -> None:
        value = value.strip()
        if value and value.lower() not in {item.lower() for item in aliases}:
            aliases.append(value)

    for item in raw_aliases:
        add(item)

    add(source_path.stem)
    add(owner)
    if owner.lower() == "duanjiabing":
        add("Duan Jiabing")
        add("段佳兵")

    for part in source_path.with_suffix("").parts:
        if contains_cjk(part):
            add(part)
            for match in re.findall(r"[\u4e00-\u9fff]{2,}", part):
                add(match)

    return aliases[:12]


def fallback_card(path: Path, context: str, doc_type: str, owner: str, topic: str) -> DocumentCard:
    title = path.stem
    summary = first_content_lines(context, 180) or f"{title} is a {doc_type} document related to {topic}."
    keyword_seed = [topic, doc_type, owner]
    keyword_seed.extend(re.split(r"[_+\\-\\s]+", path.stem))
    keywords: list[str] = []
    for item in keyword_seed:
        item = item.strip()
        if item and item.lower() not in {k.lower() for k in keywords}:
            keywords.append(item)
        if len(keywords) >= 8:
            break

    useful_for = [
        f"Finding whether this document is relevant to {topic}",
        f"Locating source material from {owner}",
    ]
    questions_answered = [
        f"{title} 主要讲了什么？",
        f"{owner} 有哪些和 {topic} 相关的资料？",
        f"哪些文档适合了解 {topic}？",
    ]
    return DocumentCard(
        schema_version="document_card.v2",
        title=title,
        aliases=merge_aliases([], path, owner),
        doc_type=doc_type,
        owner=owner,
        topic=topic,
        year=None,
        keywords=keywords,
        summary=summary,
        main_contribution="unknown",
        useful_for=useful_for,
        questions_answered=questions_answered,
        source_path=str(path),
    )


def card_to_markdown(card: DocumentCard) -> str:
    year = str(card.year) if card.year is not None else "unknown"
    keywords = ", ".join(card.keywords) if card.keywords else "unknown"
    aliases = ", ".join(card.aliases) if card.aliases else "unknown"
    useful_for = "\n".join(f"- {item}" for item in card.useful_for) or "- unknown"
    questions = "\n".join(f"- {item}" for item in card.questions_answered) or "- unknown"
    return (
        f"# {card.title}\n\n"
        "<!--\n"
        f"schema_version: {card.schema_version}\n"
        f"doc_type: {card.doc_type}\n"
        f"owner: {card.owner}\n"
        f"topic: {card.topic}\n"
        f"year: {year}\n"
        f"aliases: {aliases}\n"
        f"keywords: {keywords}\n"
        f"source_path: {card.source_path}\n"
        "-->\n\n"
        f"## Aliases\n\n{aliases}\n\n"
        f"## Summary\n\n{card.summary or 'unknown'}\n\n"
        f"## Main Contribution\n\n{card.main_contribution or 'unknown'}\n\n"
        f"## Useful For\n\n{useful_for}\n\n"
        f"## Questions Answered\n\n{questions}\n"
    )


def iter_documents(root: Path, only: str) -> list[Path]:
    files = sorted(path for path in root.rglob("*") if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS)
    if not only:
        return files
    return [path for path in files if only in path.relative_to(root).as_posix()]


def write_jsonl(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as fh:
        for row in rows:
            fh.write(json.dumps(row, ensure_ascii=False) + "\n")


def process_document(path: Path, input_root: Path, output_root: Path, args: argparse.Namespace,
                     api_key: str) -> dict[str, Any]:
    rel = path.relative_to(input_root)
    doc_dir = output_root / safe_output_slug(rel)
    json_path = doc_dir / "document_card.json"
    md_path = doc_dir / "document_card.md"
    doc_type = args.doc_type or infer_doc_type(path)
    owner = args.owner or infer_owner(path, input_root, "lab")
    topic = args.topic

    row: dict[str, Any] = {
        "source_path": str(path),
        "relative_path": rel.as_posix(),
        "output_json": str(json_path),
        "output_markdown": str(md_path),
        "doc_type": doc_type,
        "owner": owner,
        "topic": topic,
        "status": "processing",
    }
    try:
        if json_path.exists() and md_path.exists() and not args.force:
            card = coerce_card(json.loads(json_path.read_text(encoding="utf-8")), path, doc_type, owner, topic)
            row["status"] = "cached"
        else:
            context = build_model_context(path, args)
            if args.dry_run:
                print(f"\n--- {rel} extracted context preview ---\n{context[: args.preview_chars]}\n")
                row["status"] = "dry_run"
                return row
            if not api_key:
                if not args.fallback:
                    raise RuntimeError("API key is required unless --dry-run or --fallback is used")
                card = fallback_card(path, context, doc_type, owner, topic)
            else:
                try:
                    content = call_chat_completion(
                        args.endpoint,
                        api_key,
                        args.model,
                        prompt_for_card(context, str(path), doc_type, owner, topic),
                        args.temperature,
                        args.timeout,
                    )
                    card = coerce_card(parse_json_object(content), path, doc_type, owner, topic)
                    if args.enforce_english and card_needs_english_normalization(card):
                        card = normalize_card_to_english(
                            card,
                            args.endpoint,
                            api_key,
                            args.model,
                            args.temperature,
                            args.timeout,
                        )
                except Exception:
                    if not args.fallback:
                        raise
                    card = fallback_card(path, context, doc_type, owner, topic)
            doc_dir.mkdir(parents=True, exist_ok=True)
            json_path.write_text(
                json.dumps(asdict(card), ensure_ascii=False, indent=2) + "\n",
                encoding="utf-8",
            )
            md_path.write_text(card_to_markdown(card), encoding="utf-8")
            row["status"] = "ready"

        row.update(asdict(card))
    except Exception as exc:
        row["status"] = "failed"
        row["error"] = str(exc)
        if args.verbose:
            traceback.print_exc()
    return row


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Generate JSON and Markdown document cards from lab documents."
    )
    parser.add_argument("--input", default="data/tenants/RDMA", help="Source document directory")
    parser.add_argument("--output", default="data/document_cards/tenants/RDMA", help="Document-card output directory")
    parser.add_argument("--endpoint", default=DEFAULT_ENDPOINT, help="OpenAI-compatible chat completions endpoint")
    parser.add_argument("--model", default="deepseek-v4-flash", help="Model name")
    parser.add_argument("--api-key", default="", help="API key; defaults to DEEPSEEK_API_KEY")
    parser.add_argument("--only", default="", help="Only process files whose relative path contains this text")
    parser.add_argument("--limit", type=int, default=0, help="Maximum number of files to process")
    parser.add_argument("--force", action="store_true", help="Regenerate existing cards")
    parser.add_argument("--fallback", action="store_true",
                        help="Generate a local heuristic card when API key is missing or the API call fails")
    parser.add_argument("--no-enforce-english", dest="enforce_english", action="store_false",
                        help="Do not make a second normalization call when semantic fields contain Chinese")
    parser.add_argument("--dry-run", action="store_true", help="Only extract and preview text; do not call the API")
    parser.add_argument("--topic", default="RDMA")
    parser.add_argument("--owner", default="", help="Override owner for every file")
    parser.add_argument("--doc-type", default="", help="Override doc_type for every file")
    parser.add_argument("--first-pages", type=int, default=3)
    parser.add_argument("--last-pages", type=int, default=2)
    parser.add_argument("--max-slides", type=int, default=20)
    parser.add_argument("--max-paragraphs", type=int, default=80)
    parser.add_argument("--max-chars", type=int, default=18000)
    parser.add_argument("--preview-chars", type=int, default=3000)
    parser.add_argument("--temperature", type=float, default=0.1)
    parser.add_argument("--timeout", type=int, default=60)
    parser.add_argument("--verbose", action="store_true")
    parser.set_defaults(enforce_english=True)
    return parser


def main() -> int:
    args = build_parser().parse_args()
    input_root = Path(args.input).resolve()
    output_root = Path(args.output).resolve()
    api_key = args.api_key or os.environ.get("DEEPSEEK_API_KEY", "")

    if not input_root.exists():
        print(f"Input directory does not exist: {input_root}", file=sys.stderr)
        return 2

    files = iter_documents(input_root, args.only)
    if args.limit > 0:
        files = files[: args.limit]
    if not files:
        print("No supported documents matched the input filters", file=sys.stderr)
        return 2

    rows: list[dict[str, Any]] = []
    for path in files:
        row = process_document(path, input_root, output_root, args, api_key)
        rows.append(row)
        print(f"{row.get('status', 'unknown'):8s} {path.relative_to(input_root)}")

    if not args.dry_run:
        write_jsonl(output_root / "cards.jsonl", rows)
        failures = [row for row in rows if row.get("status") == "failed"]
        error_path = output_root / "errors.jsonl"
        if failures:
            write_jsonl(error_path, failures)
        elif error_path.exists():
            error_path.unlink()

    ready_count = sum(1 for row in rows if row.get("status") in {"ready", "cached"})
    failed_count = sum(1 for row in rows if row.get("status") == "failed")
    print(f"\nProcessed {len(rows)} documents: {ready_count} ready/cached, {failed_count} failed")
    if not args.dry_run:
        print(f"Cards: {output_root / 'cards.jsonl'}")
    return 1 if failed_count else 0


if __name__ == "__main__":
    raise SystemExit(main())
